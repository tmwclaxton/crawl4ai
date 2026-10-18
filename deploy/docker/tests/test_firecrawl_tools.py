"""Offline tests for the Firecrawl-shaped Crawl4ai tools."""

import asyncio
import io
import json
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

from doc_parse import ParseError, parse_document
from page_monitor import PageMonitor, deliver_webhook, parse_schedule, webhook_static_error
from site_tools import (
    ToolError,
    clamp_crawl,
    engine_cooled,
    filter_html,
    filter_mapped,
    html_to_markdown,
    is_blocked_search,
    merge_hits,
    narrow_query,
    normalize_snapshot,
    parse_relevance_answer,
    parse_search_body,
    search_hits_are_relevant,
    parse_sitemap_xml,
    path_allowed,
    rank_mapped,
    reset_search_state,
    seed_pattern,
    run_agent,
    run_map,
    run_scrape,
    apify_input,
    engine_names,
    run_search,
    search_request_url,
    seed_source,
    zip_pages,
)

FIXTURES = Path(__file__).parent / "fixtures" / "firecrawl_parity"


class MemoryRedis:
    def __init__(self):
        self.hashes = {}
        self.sets = {}

    async def hset(self, key, mapping=None, **kwargs):
        self.hashes.setdefault(key, {}).update(mapping or {})
        return 1

    async def hgetall(self, key):
        return dict(self.hashes.get(key) or {})

    async def sadd(self, key, *members):
        self.sets.setdefault(key, set()).update(members)
        return 1

    async def smembers(self, key):
        return set(self.sets.get(key) or ())

    async def srem(self, key, *members):
        bucket = self.sets.setdefault(key, set())
        for member in members:
            bucket.discard(member)
        return 1

    async def delete(self, key):
        existed = key in self.hashes
        self.hashes.pop(key, None)
        return 1 if existed else 0


@pytest.fixture(autouse=True)
def _clear_search_cooldown():
    reset_search_state()
    yield
    reset_search_state()


def _read(name: str) -> str:
    return (FIXTURES / name).read_text()


def test_fixture_page_has_the_known_sentence():
    html = _read("page.html")
    assert "enterprise plan is 42 pounds" in html
    assert "Site nav" in html


def test_scrape_main_content_include_exclude_and_query():
    html = _read("page.html")

    async def fetch(url, **kwargs):
        assert url.startswith("https://")
        if kwargs["only_main"]:
            markdown = "The enterprise plan is 42 pounds a month."
        else:
            markdown = "Site nav\nThe enterprise plan is 42 pounds a month.\nSite footer"
        return {"url": url, "markdown": markdown, "html": html, "links": ["https://example.com/pricing"]}

    async def complete(messages):
        assert "42 pounds" in messages[0]["content"]
        return "42 pounds a month"

    async def scenario():
        main = await run_scrape(
            ["https://example.com/pricing"],
            formats=["markdown", "links"],
            only_main_content=True,
            wait_for=None,
            include_tags=[],
            exclude_tags=[],
            query=None,
            fetch=fetch,
        )
        raw = await run_scrape(
            ["https://example.com/pricing"],
            formats=["markdown"],
            only_main_content=False,
            wait_for=None,
            include_tags=[],
            exclude_tags=[],
            query=None,
            fetch=fetch,
        )
        tagged = await run_scrape(
            ["https://example.com/pricing"],
            formats=["markdown"],
            only_main_content=True,
            wait_for=None,
            include_tags=["p"],
            exclude_tags=["nav", "footer"],
            query=None,
            fetch=fetch,
        )
        asked = await run_scrape(
            ["https://example.com/pricing", "https://example.com/docs"],
            formats=["markdown"],
            only_main_content=True,
            wait_for=None,
            include_tags=[],
            exclude_tags=[],
            query="What is the enterprise price?",
            fetch=fetch,
            complete=complete,
        )
        return main, raw, tagged, asked

    main, raw, tagged, asked = asyncio.run(scenario())
    assert "42 pounds" in main["data"][0]["markdown"]
    assert main["data"][0]["links"] == ["https://example.com/pricing"]
    assert "Site nav" in raw["data"][0]["markdown"]
    assert "Site nav" not in tagged["data"][0]["markdown"]
    assert "42 pounds" in tagged["data"][0]["markdown"]
    assert asked["data"][0]["answer"] == "42 pounds a month"
    assert "markdown" not in asked["data"][0]
    assert len(asked["data"]) == 2


def test_scrape_rejects_a_bad_scheme_before_fetch():
    async def fetch(url, **kwargs):
        raise AssertionError("fetch should not run")

    with pytest.raises(ToolError) as caught:
        asyncio.run(run_scrape(
            ["javascript:alert(1)"],
            formats=["markdown"],
            only_main_content=True,
            wait_for=None,
            include_tags=[],
            exclude_tags=[],
            query=None,
            fetch=fetch,
        ))
    assert caught.value.status == 400


def test_filter_html_drops_excluded_tags():
    html = filter_html(_read("page.html"), [], ["nav", "footer"])
    assert "Site nav" not in html
    assert "enterprise plan" in html


def test_html_to_markdown_keeps_the_heading():
    assert "Grant pricing" in html_to_markdown(_read("page.html"))


def test_sitemap_parse_rank_and_filters():
    urls = parse_sitemap_xml(_read("sitemap.xml"))
    assert "https://example.com/docs/auth" in urls
    items = [
        {"url": url, "head_data": {"title": "Authentication guide" if "auth" in url else "Other"}}
        for url in urls
    ]
    ranked = rank_mapped(items, "authentication")
    assert ranked[0]["url"].endswith("/docs/auth")
    kept = filter_mapped(items, root="https://example.com", include_subdomains=False, limit=20)
    assert all("blog.example.com" not in item["url"] for item in kept)
    assert all(not item["url"].endswith("robots.txt") for item in kept)
    with_sub = filter_mapped(items, root="https://example.com", include_subdomains=True, limit=20)
    assert any("blog.example.com" in item["url"] for item in with_sub)
    assert seed_source("only") == "sitemap"
    assert seed_source("skip") == "cc"
    assert seed_source("include") == "sitemap+cc"


def test_run_map_records_the_sitemap_mode_and_drops_nonsense():
    class FakeSeeder:
        def __init__(self):
            self.source = None
            self.pattern = None
            self.max_urls = None
            self.force = None

        async def urls(self, domain, config):
            self.source = config.source
            self.pattern = config.pattern
            self.max_urls = config.max_urls
            self.force = config.force
            return [
                {"url": "https://example.com/docs/auth", "head_data": {"title": "Authentication"}},
                {"url": "https://example.com/robots.txt", "head_data": {"title": "robots"}},
                {"url": "https://blog.example.com/post", "head_data": {"title": "Blog"}},
            ]

    seeder = FakeSeeder()
    result = asyncio.run(run_map(
        "https://example.com",
        limit=10,
        search="authentication",
        sitemap="only",
        include_subdomains=False,
        seeder=seeder,
    ))
    assert seeder.source == "sitemap"
    assert seeder.pattern == "*"
    assert seeder.force is True
    assert seeder.max_urls > 10
    assert result["links"] == ["https://example.com/docs/auth"]


def test_run_map_keeps_the_requested_path_and_ranks_the_search_term():
    class FakeSeeder:
        def __init__(self):
            self.pattern = None

        async def urls(self, domain, config):
            self.pattern = config.pattern
            return [
                {"url": "https://docs.python.org/3/library/functions.html"},
                {"url": "https://docs.python.org/3.11/"},
                {"url": "https://docs.python.org/3/library/constants.html"},
            ]

    seeder = FakeSeeder()
    result = asyncio.run(run_map(
        "https://docs.python.org/3/library/",
        limit=5,
        search="constants",
        sitemap="include",
        include_subdomains=False,
        seeder=seeder,
    ))
    assert seed_pattern("https://docs.python.org/3/library/") == "*docs.python.org/3/library*"
    assert seeder.pattern == "*docs.python.org/3/library*"
    assert result["links"][0].endswith("/constants.html")
    assert all("/3/library/" in link for link in result["links"])


def test_path_filters_and_crawl_clamps():
    assert path_allowed("https://example.com/docs/start", ["/docs"], [])
    assert not path_allowed("https://example.com/blog/post", ["/docs"], [])
    assert not path_allowed("https://example.com/docs/secret", ["/docs"], ["/docs/secret"])
    assert clamp_crawl(500, 9, max_pages=100, max_depth_cap=5) == (100, 5)
    with pytest.raises(ToolError):
        clamp_crawl(0, 1)


def test_search_merges_engines_skips_blocks_and_stops_at_the_limit():
    bodies = {
        "duckduckgo": _read("ddg.html"),
        "bing": _read("bing.xml"),
        "mojeek": _read("mojeek.html"),
        "searxng": _read("searx.json"),
        "brave": _read("brave.json"),
    }
    assert parse_search_body("duckduckgo", bodies["duckduckgo"])[0]["url"] == "https://example.com/docs"
    assert parse_search_body("bing", bodies["bing"])[1]["url"] == "https://bing-only.example/item"
    assert parse_search_body("mojeek", bodies["mojeek"])[0]["title"] == "Mojeek result"
    assert parse_search_body("searxng", bodies["searxng"])[0]["description"] == "From searx"
    assert is_blocked_search(200, _read("captcha.html"))
    assert is_blocked_search(403, "")
    assert not is_blocked_search(200, bodies["searxng"])

    async def fetch(name, url, headers):
        if name == "duckduckgo":
            return 200, _read("captcha.html")
        return 200, bodies[name]

    blocked = asyncio.run(run_search(
        "grant pricing",
        limit=5,
        sources="web",
        categories=None,
        apify_token="",
        fetch=fetch,
        now=1000,
    ))
    assert blocked["engines_skipped"][0] == {"engine": "duckduckgo", "reason": "blocked"}
    assert "bing" in blocked["engines_used"]
    assert engine_cooled("duckduckgo", 1001)
    urls = [hit["url"] for hit in blocked["data"]["web"]]
    assert urls.count("https://example.com/docs") == 1

    calls = []

    async def plenty(name, url, headers):
        calls.append(name)
        return 200, bodies[name]

    stopped = asyncio.run(run_search(
        "grant pricing",
        limit=2,
        sources="web",
        categories=None,
        apify_token="",
        fetch=plenty,
        now=5000,
    ))
    assert len(stopped["data"]["web"]) == 2
    assert "mojeek" not in calls

    news_url, _headers = search_request_url("bing", "grant pricing", "news")
    assert "/news/search" in news_url
    assert "site:github.com" in narrow_query("timeout error", "developer")
    assert engine_names("") == ["duckduckgo", "bing", "mojeek", "searxng"]
    assert engine_names("token")[-1] == "apify"
    assert "brave" not in engine_names("token")
    assert "memory=512" in search_request_url("apify", "grants", "web")[0]


def test_engine_fetch_error_is_skipped():
    async def boom(name, url, headers, payload=None):
        if name == "duckduckgo":
            raise TimeoutError("connect")
        if name == "bing":
            return 200, _read("bing.xml")
        return 403, ""

    result = asyncio.run(run_search(
        "grant pricing",
        limit=2,
        sources="web",
        categories=None,
        apify_token="",
        fetch=boom,
        now=40000,
    ))
    assert {"engine": "duckduckgo", "reason": "error"} in result["engines_skipped"]
    assert "bing" in result["engines_used"]
    assert result["data"]["web"]


def test_relevance_answer_is_yes_or_no():
    assert parse_relevance_answer("YES") is True
    assert parse_relevance_answer("yes, they match") is True
    assert parse_relevance_answer("NO") is False
    assert parse_relevance_answer("no - shopping results") is False
    assert parse_relevance_answer("") is None
    assert parse_relevance_answer("maybe") is None
    assert asyncio.run(search_hits_are_relevant("q", [])) is False


def test_junk_bing_hits_fall_through_to_apify():
    junk = _read("bing_junk.xml")

    async def no_match(messages):
        return "NO"

    calls = []

    async def fetch(name, url, headers, payload=None):
        calls.append(name)
        if name == "apify":
            assert "memory=512" in url
            return 201, _read("apify.json")
        if name == "bing":
            return 200, junk
        if name == "duckduckgo":
            return 200, _read("captcha.html")
        return 403, ""

    result = asyncio.run(run_search(
        "Infinitopes director linkedin",
        limit=5,
        sources="web",
        categories=None,
        apify_token="secret",
        fetch=fetch,
        now=30000,
        judge=no_match,
    ))
    assert "bing" not in result["engines_used"]
    assert {"engine": "bing", "reason": "irrelevant"} in result["engines_skipped"]
    assert "apify" in result["engines_used"]
    assert result["data"]["web"][0]["url"] == "https://docs.python.org/3/library/itertools.html"
    assert result["data"]["web"][0]["engine"] == "apify"
    assert calls[-1] == "apify"


def test_synonym_hits_kept_when_judge_says_yes():
    async def yes(messages):
        content = messages[-1]["content"]
        assert "Sarah Pinch Gingerbread" in content
        assert "linkedin.com/in/pinch" in content
        return "YES"

    async def fetch(name, url, headers, payload=None):
        if name == "duckduckgo":
            return 200, (
                '<html><body>'
                '<a class="result__a" href="https://linkedin.com/in/pinch">'
                "Sarah Pinch - Fundraising Manager at Gingerbread</a>"
                '<a class="result__snippet">UK single parent charity</a>'
                "</body></html>"
            )
        raise AssertionError(f"should not call {name}")

    result = asyncio.run(run_search(
        "Sarah Pinch Gingerbread linkedin",
        limit=3,
        sources="web",
        categories=None,
        apify_token="secret",
        fetch=fetch,
        now=31000,
        judge=yes,
    ))
    assert result["engines_used"] == ["duckduckgo"]
    assert result["data"]["web"][0]["url"] == "https://linkedin.com/in/pinch"


def test_judge_failure_keeps_engine_hits():
    async def boom_judge(messages):
        raise RuntimeError("nanogpt down")

    async def fetch(name, url, headers, payload=None):
        if name == "duckduckgo":
            raise TimeoutError("connect")
        if name == "bing":
            return 200, _read("bing.xml")
        return 403, ""

    result = asyncio.run(run_search(
        "grant pricing",
        limit=2,
        sources="web",
        categories=None,
        apify_token="",
        fetch=fetch,
        now=32000,
        judge=boom_judge,
    ))
    assert "bing" in result["engines_used"]
    assert result["data"]["web"]


def test_apify_runs_only_after_free_engines_and_parses_organic_results():
    body = _read("apify.json")
    parsed = parse_search_body("apify", body)
    assert parsed[0]["url"] == "https://docs.python.org/3/library/itertools.html"
    assert parsed[0]["engine"] == "apify"

    news = apify_input("grant deadlines", "news")
    assert news["queries"].startswith("https://www.google.com/search?tbm=nws&q=")
    assert news["maxPagesPerQuery"] == 1
    assert news["saveHtml"] is False
    assert news["websiteContentScraper"]["enable"] is False
    assert news["aiModeSearch"]["enableAiMode"] is False

    calls = []

    async def fetch(name, url, headers, payload=None):
        calls.append((name, payload))
        if name == "apify":
            assert headers["Authorization"] == "Bearer secret"
            assert "token=" not in url
            assert payload["maxPagesPerQuery"] == 1
            return 201, body
        if name == "duckduckgo":
            return 200, _read("ddg.html")
        return 403, ""

    result = asyncio.run(run_search(
        "python itertools",
        limit=5,
        sources="web",
        categories=None,
        apify_token="secret",
        fetch=fetch,
        now=9000,
    ))
    assert calls[0][0] == "duckduckgo"
    assert calls[-1][0] == "apify"
    assert "apify" in result["engines_used"]
    urls = [hit["url"] for hit in result["data"]["web"]]
    assert "https://docs.python.org/3/library/itertools.html" in urls

    filled = []

    async def plenty(name, url, headers, payload=None):
        filled.append(name)
        return 200, _read("ddg.html") if name == "duckduckgo" else _read("bing.xml")

    asyncio.run(run_search(
        "python itertools",
        limit=2,
        sources="web",
        categories=None,
        apify_token="secret",
        fetch=plenty,
        now=20000,
    ))
    assert "apify" not in filled


def test_agent_caps_pages_and_checks_the_schema():
    scraped = []

    async def scrape(url):
        scraped.append(url)
        return {"url": url, "markdown": "Ada Lovelace wrote the notes."}

    async def complete(messages):
        assert "Ada Lovelace" in messages[0]["content"]
        return '{"name":"Ada Lovelace"}'

    async def discover(prompt):
        raise AssertionError("discover should not run when urls are set")

    schema = {"type": "object", "required": ["name"], "properties": {"name": {"type": "string"}}}
    urls = [f"https://example.com/{index}" for index in range(20)]
    result = asyncio.run(run_agent(
        "Who is Ada?",
        urls,
        schema,
        scrape=scrape,
        complete=complete,
        discover=discover,
    ))
    assert result["pages_fetched"] == 8
    assert len(scraped) == 8
    assert result["data"]["name"] == "Ada Lovelace"

    async def bad_complete(messages):
        return "not json"

    with pytest.raises(ToolError) as caught:
        asyncio.run(run_agent(
            "Who?",
            ["https://example.com/a"],
            schema,
            scrape=scrape,
            complete=bad_complete,
        ))
    assert caught.value.status == 502


def _tiny_pdf() -> bytes:
    stream = b"BT /F1 18 Tf 40 80 Td (Hello from crawl4ai) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_parse_pdf_html_docx_xlsx_odt_rtf_and_rejects_legacy():
    assert "Hello from crawl4ai" in parse_document("paper.pdf", _tiny_pdf())
    assert "Grant pricing" in parse_document("page.html", _read("page.html").encode())

    from docx import Document

    document = Document()
    document.add_paragraph("Hello from docx")
    buffer = io.BytesIO()
    document.save(buffer)
    assert "Hello from docx" in parse_document("notes.docx", buffer.getvalue())

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet["A1"] = "enterprise"
    sheet["B1"] = 42
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    rendered = parse_document("prices.xlsx", xlsx.getvalue())
    assert "enterprise" in rendered and "42" in rendered

    odt_xml = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        b'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        b"<office:body><office:text><text:p>Hello from odt</text:p></office:text></office:body>"
        b"</office:document-content>"
    )
    odt = io.BytesIO()
    with zipfile.ZipFile(odt, "w") as archive:
        archive.writestr("content.xml", odt_xml)
    assert "Hello from odt" in parse_document("notes.odt", odt.getvalue())
    assert "Hello from rtf" in parse_document("notes.rtf", br"{\rtf1\ansi Hello from rtf}")

    with pytest.raises(ParseError) as legacy:
        parse_document("old.doc", b"not a real doc")
    assert legacy.value.status == 400
    with pytest.raises(ParseError):
        parse_document("old.xls", b"not a real xls")
    with pytest.raises(ParseError) as huge:
        parse_document("big.pdf", b"x" * (8 * 1024 * 1024 + 1))
    assert huge.value.status == 413


def test_download_zip_layout_and_page_cap():
    pages = [
        {"url": "https://example.com/docs/start", "markdown": "# Start", "links": ["https://example.com/docs/next"]},
        {"url": "https://example.com/", "markdown": "# Home", "links": []},
    ]
    blob = zip_pages(pages, ["markdown", "links"])
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        names = set(archive.namelist())
        assert "docs/start.md" in names
        assert "docs/start.links.txt" in names
        assert archive.read("docs/start.md") == b"# Start"
    too_many = [{"url": f"https://example.com/{index}", "markdown": "x", "links": []} for index in range(21)]
    with pytest.raises(ToolError):
        zip_pages(too_many, ["markdown"])


def test_monitor_labels_goal_webhook_and_delete():
    redis = MemoryRedis()
    pages = {
        "https://example.com/pricing": "The enterprise plan is 19 pounds. https://example.com/?id=1&utm_source=ad"
    }
    delivered = []

    async def fetch(url):
        if url.endswith("/missing"):
            return None
        if url.endswith("/boom"):
            raise RuntimeError("network down")
        return pages[url]

    async def judge(goal, before, after):
        assert "price" in goal
        return "24" in after

    async def notify(url, payload):
        delivered.append((url, payload))

    monitor = PageMonitor(redis, fetch=fetch, judge=judge, notify=notify)

    async def scenario():
        with pytest.raises(ToolError):
            await monitor.create(
                name="Too fast",
                urls=["https://example.com/pricing"],
                schedule="every 5 minutes",
                goal=None,
                webhook=None,
            )
        assert webhook_static_error("http://example.com/hook")
        assert webhook_static_error("https://127.0.0.1/hook")
        assert webhook_static_error("https://localhost/hook")
        assert webhook_static_error("https://10.1.1.1/hook")
        created = await monitor.create(
            name="Pricing",
            urls=["https://example.com/pricing", "https://example.com/missing", "https://example.com/boom"],
            schedule="hourly",
            goal="Alert when the price changes.",
            webhook="https://alerts.example/hook",
        )
        first = await monitor.check(created["id"])
        statuses = {page["url"]: page["status"] for page in first["pages"]}
        second = await monitor.check(created["id"])
        pages["https://example.com/pricing"] = "The enterprise plan is    19 pounds. https://example.com/?utm_source=other&id=1"
        tracking = await monitor.check(created["id"])
        pages["https://example.com/pricing"] = "The enterprise plan is 24 pounds."
        changed = await monitor.check(created["id"])
        listed = await monitor.list_monitors()
        deleted = await monitor.delete(created["id"])
        with pytest.raises(ToolError) as missing:
            await monitor.delete(created["id"])
        return statuses, second, tracking, changed, listed, deleted, missing

    statuses, second, tracking, changed, listed, deleted, missing = asyncio.run(scenario())
    assert statuses["https://example.com/pricing"] == "new"
    assert statuses["https://example.com/missing"] == "removed"
    assert statuses["https://example.com/boom"] == "error"
    assert second["pages"][0]["status"] == "same"
    assert tracking["pages"][0]["status"] == "same"
    assert changed["pages"][0]["status"] == "changed"
    assert delivered[-1][0] == "https://alerts.example/hook"
    assert listed["monitors"][0]["name"] == "Pricing"
    assert deleted["deleted"]
    assert missing.value.status == 404
    assert parse_schedule("every 30 minutes") == 1800
    assert parse_schedule("daily at 9:00") == 86400


def test_goal_judge_can_ignore_a_real_looking_edit():
    redis = MemoryRedis()
    state = {"text": "Price is 19 today."}

    async def fetch(url):
        return state["text"]

    async def judge(goal, before, after):
        return False

    monitor = PageMonitor(redis, fetch=fetch, judge=judge)

    async def scenario():
        created = await monitor.create(
            name="Quiet",
            urls=["https://example.com/pricing"],
            schedule="every 15 minutes",
            goal="Alert on price changes.",
            webhook=None,
        )
        await monitor.check(created["id"])
        state["text"] = "Price is 19 today. Regenerated footer."
        return await monitor.check(created["id"])

    result = asyncio.run(scenario())
    assert result["pages"][0]["status"] == "same"


def test_webhook_reaches_a_local_catcher():
    received = {}

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            length = int(self.headers.get("Content-Length", "0"))
            received["body"] = self.rfile.read(length)
            self.send_response(204)
            self.end_headers()

        def log_message(self, format, *args):
            return

    server = HTTPServer(("127.0.0.1", 0), Handler)
    port = server.server_address[1]
    thread = threading.Thread(target=server.handle_request)
    thread.start()
    try:
        asyncio.run(deliver_webhook(
            f"http://127.0.0.1:{port}/hook",
            {"pages": [{"status": "changed"}]},
        ))
    finally:
        thread.join(timeout=5)
        server.server_close()
    assert json.loads(received["body"])["pages"][0]["status"] == "changed"


def test_merge_hits_dedupes_and_stops():
    groups = [
        [{"url": "https://a.example", "title": "A"}],
        [{"url": "https://a.example", "title": "A again"}, {"url": "https://b.example", "title": "B"}],
    ]
    assert [hit["url"] for hit in merge_hits(groups, 1)] == ["https://a.example"]


def test_normalize_snapshot_drops_tracking_and_whitespace():
    left = normalize_snapshot("Hello   WORLD https://example.com/a?utm_source=x&id=1")
    right = normalize_snapshot("hello world https://example.com/a?id=1")
    assert left == right
